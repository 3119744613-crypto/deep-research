"""Service acceptance with synthetic ordinary document parsing requests."""

import asyncio
import json
import stat
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from document_research.config import init_config, load_config, validate_data_dir
from document_research.contracts import ResearchError
from document_research.providers import ProviderError
from document_research.service import PREFIX, create_app
from document_research.storage import Store

ALICE = "synthetic-alice-token-at-least-24-chars"
BOB = "synthetic-bob-token-at-least-24-chars"


class SyntheticEngine:
    def __init__(self):
        self.release = threading.Event()
        self.release.set()
        self.started = threading.Event()
        self.configured = True
        self.seen = []
        self.failure = None

    def is_configured(self):
        return self.configured

    async def run(self, task, sources, authorize):
        authorize()
        self.seen.append((task, sources))
        self.started.set()
        while not self.release.is_set():
            await asyncio.sleep(0.01)
        authorize()
        if self.failure:
            if isinstance(self.failure, Exception):
                raise self.failure
            raise RuntimeError(self.failure)
        return {
            "summary": "合成测试报告，非真实实验",
            "assumptions": [],
            "sources": sources,
            "measurement_status": "not_measured",
        }


def configure(path, **changes):
    config_path = path / "config.json"
    config = json.loads(config_path.read_text())
    config.update(changes)
    config_path.write_text(json.dumps(config))
    return config


@pytest.fixture
def setup(tmp_path):
    init_config(tmp_path)
    configure(
        tmp_path,
        tokens={"alice": ALICE, "bob": BOB},
        zhipu={
            "api_key": "synthetic-never-real",
            "libraries": {
                "alice": [{"id": "docs-one", "name": "只读解析文档"}],
                "bob": [],
            },
        },
    )
    engine = SyntheticEngine()
    app = create_app(tmp_path, engine)
    with TestClient(app, headers={"Authorization": "Bearer " + ALICE}) as client:
        yield client, tmp_path, engine, app


def library(client, name="表格需求"):
    response = client.post(PREFIX + "/libraries", json={"name": name})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def upload(
    client,
    library_id,
    name="requirements.md",
    content=b"Merged table cells fail. Expected all headers.",
):
    return client.post(
        f"{PREFIX}/libraries/{library_id}/documents", files={"file": (name, content)}
    )


def task(client, library_ids=None, **changes):
    body = {
        "problem": "复杂表格解析漏掉合并单元格，请提出对比实验。",
        "library_ids": library_ids or [],
        "parser_version": "synthetic-parser 1.0",
        "expected_result": "保留表头与合并单元格",
        "constraints": "CPU、本地测试",
        "allow_external_model": True,
    }
    body.update(changes)
    response = client.post(PREFIX + "/tasks", json=body)
    assert response.status_code == 201, response.text
    return response.json()


def run(client, item):
    return client.post(f"{PREFIX}/tasks/{item['id']}/run", json={})


def wait_status(client, item, expected="completed"):
    for _ in range(100):
        response = client.get(f"{PREFIX}/tasks/{item['id']}")
        if response.status_code == 200 and response.json()["status"] == expected:
            return response.json()
        time.sleep(0.01)
    pytest.fail(
        f"task did not reach {expected}: {response.status_code} {response.text}"
    )


def test_private_configuration_and_separate_directory(tmp_path):
    config = init_config(tmp_path)
    assert stat.S_IMODE(tmp_path.stat().st_mode) == 0o700
    assert stat.S_IMODE(config.stat().st_mode) == 0o600
    assert len(load_config(tmp_path)["tokens"]["local"]) > 24
    with pytest.raises(FileExistsError):
        init_config(tmp_path)
    with pytest.raises(ValueError):
        validate_data_dir("relative")
    with pytest.raises(ValueError):
        validate_data_dir(Path(__file__).resolve().parents[1] / "data")
    with pytest.raises(ValueError):
        validate_data_dir(tmp_path / "backend" / "data")
    config.chmod(0o644)
    with pytest.raises(ValueError, match="600"):
        load_config(tmp_path)


def test_health_session_and_current_token_revocation(setup):
    client, path, _, _ = setup
    assert (
        client.get(PREFIX + "/health", headers={"Authorization": ""}).status_code == 200
    )
    assert client.get(PREFIX + "/session").json() == {"user": "alice"}
    assert (
        client.get(PREFIX + "/libraries", headers={"Authorization": ""}).status_code
        == 401
    )
    configure(path, tokens={"bob": BOB})
    assert client.get(PREFIX + "/libraries").status_code == 401


def test_local_upload_preview_and_owner_isolation(setup):
    client, path, _, _ = setup
    lib = library(client)
    response = upload(client, lib, "../../sample.md")
    assert response.status_code == 201, response.text
    item = response.json()
    assert item["name"] == "sample.md"
    assert item["status"] == "ready"
    preview = client.get(f"{PREFIX}/documents/{item['id']}").json()
    assert preview["chunks"][0]["page"] is None
    assert "Merged table" in preview["chunks"][0]["text"]
    assert stat.S_IMODE((path / "document-research.sqlite3").stat().st_mode) == 0o600
    assert not (path / "sample.md").exists()
    bob = {"Authorization": "Bearer " + BOB}
    assert (
        client.get(f"{PREFIX}/documents/{item['id']}", headers=bob).status_code == 403
    )
    assert (
        client.post(
            PREFIX + "/tasks",
            json={"problem": "compare parsers", "library_ids": [lib]},
            headers=bob,
        ).status_code
        == 403
    )
    assert (
        client.delete(f"{PREFIX}/documents/{item['id']}", headers=bob).status_code
        == 403
    )
    assert client.delete(f"{PREFIX}/documents/{item['id']}").status_code == 200
    assert client.get(f"{PREFIX}/documents/{item['id']}").status_code == 404


def test_readonly_zhipu_grants_are_operator_configured(setup):
    client, path, _, _ = setup
    items = client.get(PREFIX + "/libraries").json()["items"]
    assert items == [
        {
            "id": "zhipu:docs-one",
            "name": "只读解析文档",
            "kind": "zhipu",
            "document_count": 0,
        }
    ]
    assert client.delete(PREFIX + "/libraries/zhipu:docs-one").status_code == 400
    assert upload(client, "zhipu:docs-one").status_code == 400
    assert "api_key" not in client.get(PREFIX + "/libraries").text
    configure(path, zhipu={"libraries": {}})
    assert client.get(PREFIX + "/libraries").json()["items"] == []


def test_failed_files_preserved_scanned_pdf_and_zip_bomb(setup):
    from pypdf import PdfWriter

    client, _, _, _ = setup
    lib = library(client)
    pdf = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(pdf)
    scanned = upload(client, lib, "scan.pdf", pdf.getvalue()).json()
    assert scanned["status"] == "failed" and "OCR" in scanned["error"]
    archive = BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as file:
        file.writestr("word/document.xml", "X" * 2_000_000)
    bomb = upload(client, lib, "bomb.docx", archive.getvalue()).json()
    assert bomb["status"] == "failed" and "压缩" in bomb["error"]
    invalid = upload(client, lib, "file.exe", b"invalid").json()
    assert invalid["status"] == "failed"
    assert len(client.get(f"{PREFIX}/libraries/{lib}/documents").json()["items"]) == 3


def test_docx_table_text_has_no_fabricated_page(setup):
    from docx import Document

    client, _, _, _ = setup
    doc = Document()
    doc.add_paragraph("要求保留复杂表头")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "方案"
    table.cell(0, 1).text = "约束"
    data = BytesIO()
    doc.save(data)
    result = upload(
        client, library(client), "requirements.docx", data.getvalue()
    ).json()
    assert result["status"] == "ready", result
    assert result["chunks"][0]["page"] is None
    assert "方案 | 约束" in result["chunks"][0]["text"]


def test_pdf_text_preserves_original_page_number(setup):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    client, _, _, _ = setup
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    page = writer.add_blank_page(width=100, height=100)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 30 Td (Expected table headers) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    data = BytesIO()
    writer.write(data)
    result = upload(client, library(client), "requirements.pdf", data.getvalue()).json()
    assert result["status"] == "ready", result
    assert result["chunks"][0]["page"] == 2
    assert result["chunks"][0]["text"] == "Expected table headers"


def test_upload_limit_rejects_before_persisting(setup):
    client, _, _, _ = setup
    lib = library(client)
    response = upload(client, lib, "large.txt", b"x" * (10 * 1024 * 1024 + 1))
    assert response.status_code == 413, response.text
    assert client.get(f"{PREFIX}/libraries/{lib}/documents").json()["items"] == []
    response = upload(client, lib, "large.txt", b"x" * (11 * 1024 * 1024))
    assert response.status_code == 413, response.text


def test_three_questions_assumptions_and_no_fabricated_measurement(setup):
    client, _, _, _ = setup
    item = task(client, parser_version="", expected_result="", constraints="")
    assert item["status"] == "needs_input"
    assert len(item["questions"]) == 3
    assert run(client, item).status_code == 409
    patched = client.patch(
        f"{PREFIX}/tasks/{item['id']}", json={"accept_assumptions": True}
    ).json()
    assert patched["status"] == "ready"
    assert run(client, item).status_code == 202
    complete = wait_status(client, item)
    assert len(complete["report"]["assumptions"]) == 3
    assert complete["report"]["measurement_status"] == "not_measured"


def test_no_model_or_no_consent_returns_real_failure(setup):
    client, _, engine, _ = setup
    no_consent = task(client, allow_external_model=False)
    assert run(client, no_consent).status_code == 400
    engine.configured = False
    no_model = task(client)
    assert run(client, no_model).status_code == 503
    assert not engine.seen


def test_only_selected_sources_and_empty_selection(setup):
    client, _, engine, _ = setup
    selected = library(client, "selected")
    other = library(client, "not selected")
    upload(client, selected, content=b"allowed evidence")
    upload(client, other, content=b"NEVER INCLUDE")
    item = task(client, [selected])
    assert run(client, item).status_code == 202
    complete = wait_status(client, item)
    assert len(complete["report"]["sources"]) == 1
    assert complete["report"]["sources"][0]["text"] == "allowed evidence"
    empty = task(client)
    assert run(client, empty).status_code == 202
    wait_status(client, empty)
    assert engine.seen[-1][1] == []
    assert (
        client.get(
            f"{PREFIX}/tasks/{item['id']}", headers={"Authorization": "Bearer " + BOB}
        ).status_code
        == 404
    )


def test_cancel_and_duplicate_run_preserve_cancelled_status(setup):
    client, _, engine, _ = setup
    engine.release.clear()
    item = task(client)
    assert run(client, item).status_code == 202
    assert engine.started.wait(1)
    assert run(client, item).status_code == 409
    assert (
        client.patch(
            f"{PREFIX}/tasks/{item['id']}", json={"constraints": "new"}
        ).status_code
        == 409
    )
    response = client.post(f"{PREFIX}/tasks/{item['id']}/cancel", json={})
    assert response.json()["status"] == "cancelled"
    engine.release.set()
    assert wait_status(client, item, "cancelled")["report"] is None


@pytest.mark.parametrize("revoke", ["token", "remote_grant", "document", "library"])
def test_midrun_revocation_and_deletion_discard_result(setup, revoke):
    client, path, engine, app = setup
    engine.release.clear()
    lib = library(client)
    document = upload(client, lib).json()
    item = task(client, [lib, "zhipu:docs-one"])
    assert run(client, item).status_code == 202
    assert engine.started.wait(1)
    if revoke == "token":
        configure(path, tokens={"bob": BOB})
    elif revoke == "remote_grant":
        configure(path, zhipu={"libraries": {}})
    elif revoke == "document":
        assert client.delete(f"{PREFIX}/documents/{document['id']}").status_code == 200
    else:
        assert client.delete(f"{PREFIX}/libraries/{lib}").status_code == 200
    engine.release.set()
    for _ in range(100):
        stored = app.state.store.task(item["id"])
        if stored["status"] == "failed":
            break
        time.sleep(0.01)
    assert stored["status"] == "failed"
    assert stored["report"] is None
    assert client.get(f"{PREFIX}/tasks/{item['id']}").status_code in (401, 403)


def test_completed_report_becomes_inaccessible_after_document_delete(setup):
    client, _, _, _ = setup
    lib = library(client)
    document = upload(client, lib).json()
    item = task(client, [lib])
    run(client, item)
    wait_status(client, item)
    client.delete(f"{PREFIX}/documents/{document['id']}")
    assert client.get(f"{PREFIX}/tasks/{item['id']}").status_code == 403


def test_restart_marks_running_task_failed_and_redacts_provider_error(setup):
    client, path, engine, app = setup
    item = task(client)
    stored = app.state.store.task(item["id"])
    stored["status"] = "running"
    app.state.store.save_task(stored)
    fresh_store = Store(path)
    fresh_store.fail_interrupted()
    assert fresh_store.task(item["id"])["status"] == "failed"
    engine.failure = "api_key=VERY_SECRET_PROVIDER_KEY"
    run(client, item)
    result = wait_status(client, item, "failed")
    assert "VERY_SECRET" not in json.dumps(result)
    assert result["report"] is None


def test_invalid_input_has_uniform_safe_error(setup):
    client, _, _, _ = setup
    response = client.post(
        PREFIX + "/tasks", json={"problem": "", "token": "private-input"}
    )
    assert response.status_code == 422
    assert isinstance(response.json()["detail"], str)
    assert "private-input" not in response.text


@pytest.mark.parametrize(
    "error",
    [ResearchError("证据引用无效，请重新生成。"), ProviderError("模型请求超时。")],
)
def test_known_safe_failures_remain_actionable(setup, error):
    client, _, engine, _ = setup
    engine.failure = error
    item = task(client)
    run(client, item)
    result = wait_status(client, item, "failed")
    assert result["error"] == str(error)


def test_second_service_cannot_share_live_data_directory(setup):
    _, path, _, _ = setup
    second = create_app(path, SyntheticEngine())
    with pytest.raises(RuntimeError, match="已有服务"), TestClient(second):
        pass


def test_atomic_completion_cannot_overwrite_cancellation_or_changed_sources(setup):
    client, _, _, app = setup
    lib = library(client)
    item = task(client, [lib])
    stored = app.state.store.task(item["id"])
    stored["status"] = "cancelled"
    app.state.store.save_task(stored)
    completed = dict(stored, status="completed", report={"private": "evidence"})
    assert app.state.store.finish_running(completed) is False
    assert app.state.store.task(item["id"])["status"] == "cancelled"
    stored["status"] = "running"
    app.state.store.save_task(stored)
    upload(client, lib)
    with pytest.raises(PermissionError):
        app.state.store.finish_running(completed, {lib: 0})
    assert app.state.store.task(item["id"])["report"] is None


def test_concurrent_patch_cannot_replace_started_or_completed_run(setup, monkeypatch):
    """交错两个真实 HTTP 请求，复现线程池 PATCH 的过期快照覆盖研究结果。"""
    client, _, _, app = setup
    item = task(client)
    read = app.state.store.task
    patch_has_snapshot = threading.Event()
    release_patch = threading.Event()
    first_read = True

    def delayed_read(task_id):
        nonlocal first_read
        snapshot = read(task_id)
        if first_read:
            first_read = False
            patch_has_snapshot.set()
            assert release_patch.wait(5), "PATCH synchronization timed out"
        return snapshot

    monkeypatch.setattr(app.state.store, "task", delayed_read)
    with ThreadPoolExecutor(max_workers=1) as executor:
        pending_patch = executor.submit(
            client.patch,
            f"{PREFIX}/tasks/{item['id']}",
            json={"constraints": "edited constraints"},
        )
        try:
            assert patch_has_snapshot.wait(2)
            assert run(client, item).status_code == 202
            completed = wait_status(client, item)
        finally:
            release_patch.set()
        patch_response = pending_patch.result(timeout=2)
    assert patch_response.status_code == 409, patch_response.text
    final = client.get(f"{PREFIX}/tasks/{item['id']}").json()
    assert final["status"] == "completed"
    assert final["report"] == completed["report"]
    assert final["constraints"] == item["constraints"]
